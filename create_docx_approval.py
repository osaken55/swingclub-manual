import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

GDRIVE_DIR = "/Users/junshu/Library/CloudStorage/GoogleDrive-osaken@gmail.com/マイドライブ/【TSH公式提案】SwingClub次世代利用者ポータル＆AIパック"
OUTPUT_DOCX = os.path.join(GDRIVE_DIR, "02_【社内稟議用】事業計画書・投資対効果ROI（Googleドキュメント版）.docx")

doc = Document()

# Page Margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
PRIMARY = RGBColor(16, 185, 129)     # #10B981
DARK_BLUE = RGBColor(15, 23, 42)     # #0F172A
GRAY = RGBColor(100, 116, 139)       # #64748B

# Header / Title
p_top = doc.add_paragraph()
p_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r_date = p_top.add_run("起案日：2026年8月吉日\n起案部署：ゴルフシステムサービス部\n起案者：部長 北村 悟志\n決裁区分：役員会 / 経営会議")
r_date.font.size = Pt(10)
r_date.font.color.rgb = GRAY
r_date.font.name = "Hiragino Sans"

title_p = doc.add_paragraph()
title_run = title_p.add_run("【社内稟議書・事業計画書】\n『SwingClub-CLOUD』ユーザー向け実践Webポータル・紙芝居動画・AIナレッジシステム 導入および外注制作の件")
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = DARK_BLUE
title_run.font.name = "Hiragino Sans"

doc.add_paragraph("─" * 45)

# Section 1
h1 = doc.add_heading("1. 稟議件名", level=1)
p1 = doc.add_paragraph("『SwingClub-CLOUD ユーザー向け実践Webポータル・スライド動画・AIナレッジシステム』外注制作および全国200コース展開の件")
p1.runs[0].font.size = Pt(11)

# Section 2
doc.add_heading("2. 目的と導入の背景", level=1)

p2_1 = doc.add_paragraph()
r = p2_1.add_run("① サポートデスクの逼迫と初歩的問い合わせの削減\n")
r.bold = True
p2_1.add_run("現在、全国約200コースのゴルフ場ユーザーから、日々「予約の入れ方」「同伴者の入力ルール」「コンペ集計の手順」「資格と料金の連動」等に関する初歩的な操作問い合わせがTSHサポートデスクに集中しております。現行マニュアルは機能仕様書型であり、現場スタッフが「業務シナリオ」から逆引きできないため、自己解決できずサポート電話に依存する構造になっております。")

p2_2 = doc.add_paragraph()
r = p2_2.add_run("② 競合優位性の確立とクラウド版への移行促進\n")
r.bold = True
p2_2.add_run("クラウド型基幹システム市場において、「現場スタッフがスマホで1分で学べる紙芝居動画」「自然言語で即答するAIヘルプデスク」を公式標準装備することで、競合他社に対する決定的な差別化要因となります。新規提案時にも支配人・役員層に対して「導入後の教育コストがゼロになる」という強力なキラーコンテンツとなります。")

# Section 3: Table
doc.add_heading("3. 投資対効果（ROI 試算）", level=1)

table = doc.add_table(rows=4, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

headers = ["項目", "現状の課題・コスト", "導入後の効果（年間試算）"]
hdr_cells = table.rows[0].cells
for i, h in enumerate(headers):
    hdr_cells[i].text = h
    hdr_cells[i].paragraphs[0].runs[0].font.bold = True
    hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    shading = parse_xml(r'<w:shd {} w:fill="0F172A"/>'.format(nsdecls('w')))
    hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

data = [
    ("サポートデスク人件費", "月間約300件の操作質問対応（月75時間）", "問い合わせ件数 約40%削減\n（年間約360時間削減 ➔ 約180万円相当/年のコスト抑制）"),
    ("新規導入時の現地説明会", "1コースあたり2〜3日の現地講習・出張拘束", "事前動画・Web学習により現地講習を1日に短縮\n（年間20コース導入で 約200万円の出張・人件費削減）"),
    ("営業成約率・リプレイス", "他社基幹との機能差が僅差化", "「AIマニュアル・動画教材つき」による受注増\n（年間1〜2件獲得増 ➔ 約500万〜1,000万円の売上増）")
]

for row_idx, row_data in enumerate(data):
    row_cells = table.rows[row_idx + 1].cells
    for col_idx, text in enumerate(row_data):
        row_cells[col_idx].text = text
        shading_color = "F1F5F9" if row_idx % 2 == 1 else "FFFFFF"
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{shading_color}"/>')
        row_cells[col_idx]._tc.get_or_add_tcPr().append(shd)

p_roi_sum = doc.add_paragraph()
r = p_roi_sum.add_run("\n➔ 結論：年間約380万円以上のコスト削減効果が見込まれ、単年度で投下資本（350万円）を全額回収可能です。")
r.bold = True
r.font.color.rgb = PRIMARY

# Section 4
doc.add_heading("4. 委託先および選定理由", level=1)
doc.add_paragraph("・委託先：津カントリー倶楽部 DXプロジェクト（代表：社長顧問 長田 賢一郎 氏）\n"
                  "・選定理由：\n"
                  "  1. SwingClubの実機運用・ゴルフ場現場業務（フロント・マスター室・経理）を完全に熟知。\n"
                  "  2. 生成AI（Gemini 2.5 Flash / NotebookLM）およびスライド動画制作の高度な内製パイプラインを保有。\n"
                  "  3. 大手ITマニュアル制作会社（見積り800万円超）と比較し、約半額での一括制作受託が可能。")

# Section 5
doc.add_heading("5. 制作内容および契約金額", level=1)
doc.add_paragraph("・契約金額（一括委託）：3,500,000 円（税別）※標準竹プラン\n"
                  "・制作期間：約2〜3ヶ月\n"
                  "・納品物：\n"
                  "  ① 利用者目線 実践マニュアル（全20モジュール：HTML / Markdown）\n"
                  "  ② レスポンシブWebポータル（静的サイト一式）\n"
                  "  ③ 実践「紙芝居」スライド動画（全10話・1080p MP4）\n"
                  "  ④ AIナレッジパッケージ（Google NotebookLM / Gemini連携データ）")

doc.save(OUTPUT_DOCX)
print(f"Document saved successfully to: {OUTPUT_DOCX}")
