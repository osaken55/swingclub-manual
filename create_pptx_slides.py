import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

GDRIVE_DIR = "/Users/junshu/Library/CloudStorage/GoogleDrive-osaken@gmail.com/マイドライブ/【TSH公式提案】SwingClub次世代利用者ポータル＆AIパック"
OUTPUT_PPTX = os.path.join(GDRIVE_DIR, "01_【プレゼン】事業提案スライド（Googleスライド版）.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
BG_DARK = RGBColor(15, 23, 42)       # #0F172A
BG_CARD = RGBColor(30, 41, 59)       # #1E293B
PRIMARY = RGBColor(16, 185, 129)     # #10B981 (Emerald)
PRIMARY_DARK = RGBColor(4, 120, 87)  # #047857
TEXT_WHITE = RGBColor(255, 255, 255)
TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
ACCENT = RGBColor(245, 158, 11)      # #F59E0B
DANGER = RGBColor(239, 68, 68)       # #EF4444

def set_slide_background(slide, color=BG_DARK):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category="SwingClub-CLOUD 次世代利用者ポータル＆実践AIナレッジ"):
    # Header bar
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.0))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_CARD
    shape.line.fill.background()
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = category
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED
    p.font.name = "Hiragino Sans"
    
    # Title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.8))
    p2 = tb.text_frame.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY
    p2.font.name = "Hiragino Sans"

# Slide 1: Cover
s1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s1)

# Badge
badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(3.2), Inches(0.45))
badge.fill.solid()
badge.fill.fore_color.rgb = PRIMARY_DARK
badge.line.fill.background()
badge.text_frame.paragraphs[0].text = "TSH公式 提案パッケージ"
badge.text_frame.paragraphs[0].font.size = Pt(14)
badge.text_frame.paragraphs[0].font.bold = True
badge.text_frame.paragraphs[0].font.color.rgb = TEXT_WHITE
badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# Main Title
tb = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(2.2))
p = tb.text_frame.paragraphs[0]
p.text = "SwingClub-CLOUD\n次世代利用者ポータル ＆ 実践AIナレッジ"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = TEXT_WHITE
p.font.name = "Hiragino Sans"

# Subtitle
tb2 = s1.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(11.3), Inches(1.2))
p = tb2.text_frame.paragraphs[0]
p.text = "ゴルフ場現場の自立運用と、TSHサポート工数半減を実現する新・公式マニュアル体系"
p.font.size = Pt(22)
p.font.color.rgb = PRIMARY
p.font.name = "Hiragino Sans"

# Footer
tb3 = s1.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.8))
p = tb3.text_frame.paragraphs[0]
p.text = "津カントリー倶楽部 DXプロジェクト | 社長顧問 長田 賢一郎\n提案先: 東京システムハウス株式会社 ゴルフシステムサービス部 部長 北村 悟志 様"
p.font.size = Pt(15)
p.font.color.rgb = TEXT_MUTED
p.font.name = "Hiragino Sans"

# Slide 2: Problem
s2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s2)
add_header(s2, "現場が抱える「見えない課題」")

issues = [
    ("1. マニュアルの孤立", "基幹サーバー内やExcel/紙に閉じ込められ、フロントやコース現場ですぐに開けない。"),
    ("2. 開発者目線の壁", "ボタン機能の説明ばかりで、「電話受付」「当日同伴者変更」等の現場シナリオが見えない。"),
    ("3. 新人教育の属人化", "先輩の口伝に依存し、誤った操作（資格コピー、同伴者名複製等）が現場で再発する。")
]

for i, (title, desc) in enumerate(issues):
    x = Inches(1.0 + i * 3.8)
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.3), Inches(3.6), Inches(3.2))
    card.fill.solid()
    card.fill.fore_color.rgb = BG_CARD
    card.line.fill.background()
    
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"\n{title}\n\n"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_WHITE

# Alert Bottom
alert = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(5.8), Inches(11.3), Inches(0.9))
alert.fill.solid()
alert.fill.fore_color.rgb = RGBColor(63, 29, 36)
alert.line.color.rgb = DANGER
alert.line.width = Pt(2)
p = alert.text_frame.paragraphs[0]
p.text = "🚨 結果：TSHサポートデスクへ「初歩的な操作の質問電話」が集中し、貴重なエンジニア稼働を圧迫！"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(254, 202, 202)
p.alignment = PP_ALIGN.CENTER

# Slide 3: Solution
s3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s3)
add_header(s3, "解決策（To-Be）: 読まないマニュアルから引けるポータルへ")

sols = [
    ("📱 ① レスポンシブWebポータル", "スマホ・タブレット・PC全対応。\nいつでも10秒で逆引き検索できるモダンUI。"),
    ("🎬 ② 実践「紙芝居」動画（全10話）", "1話90秒の短尺スライド動画。\n画面スクショと赤丸でNG操作を直感理解。"),
    ("🤖 ③ 自然言語AI相談室（Gemini）", "NotebookLM / Gemini 2.5 Flash連携。\n「同伴者未定の時は？」と聞くだけで即答。")
]

for i, (title, desc) in enumerate(sols):
    x = Inches(1.0 + i * 3.8)
    card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.4), Inches(3.6), Inches(3.8))
    card.fill.solid()
    card.fill.fore_color.rgb = BG_CARD
    card.line.color.rgb = PRIMARY
    card.line.width = Pt(2)
    
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"\n{title}\n\n"
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_WHITE

# Slide 4: Demo / Proof
s4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s4)
add_header(s4, "実証デモ：TCC先行稼働実績とYouTube紙芝居動画")

card1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.3), Inches(5.4), Inches(4.3))
card1.fill.solid()
card1.fill.fore_color.rgb = BG_CARD
card1.line.fill.background()
tf1 = card1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "🎬 第2話 紙芝居動画（74秒実演）\n"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = ACCENT
p2 = tf1.add_paragraph()
p2.text = "・タイトル：『もう迷わない！予約受付と半角カナ検索の魔法』\n・YouTube限定公開 URL：\n  https://youtu.be/wLv5pd2OoP0\n・実機スクショ＋赤丸注記＋クリアナレーションで構成。"
p2.font.size = Pt(15)
p2.font.color.rgb = TEXT_WHITE

card2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(2.3), Inches(5.4), Inches(4.3))
card2.fill.solid()
card2.fill.fore_color.rgb = BG_CARD
card2.line.fill.background()
tf2 = card2.text_frame
tf2.word_wrap = True
p = tf2.paragraphs[0]
p.text = "🧠 NotebookLM / Web実機体験\n"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = PRIMARY
p2 = tf2.add_paragraph()
p2.text = "・オサケンさん構築のSwingClub専用NotebookLMを共有\n・スマホから自然言語・音声で質問可能\n・Webポータル（index.html）で右下のGeminiチャットが即座に回答するデモを実装済み。"
p2.font.size = Pt(15)
p2.font.color.rgb = TEXT_WHITE

# Slide 5: Benefits
s5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s5)
add_header(s5, "TSH様が得られる「4大メリット」")

bens = [
    ("1. サポート問い合わせ件数の激減（約40%削減）", "現場の自己解決が進み、サポート一次対応コストとエンジニアの拘束時間を大幅削減。"),
    ("2. 競合基幹システムに対する圧倒的な優位性", "「導入後、誰でも1日で覚えられる直感ポータル付き」として他社システムに大差をつける。"),
    ("3. 営業提案・リプレイス時のキラーコンテンツ", "ゴルフ場支配人・役員層へ「導入後のスタッフ教育コストがゼロになる」と強力訴求。"),
    ("4. 新規導入時の現地オンボーディング工数半減", "事前Web・動画学習により、現地講習日数を2〜3日から1日に短縮。出張コスト削減。")
]

for i, (title, desc) in enumerate(bens):
    y = Inches(2.2 + i * 1.15)
    box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y, Inches(11.3), Inches(0.95))
    box.fill.solid()
    box.fill.fore_color.rgb = BG_CARD
    box.line.fill.background()
    
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"{title}: "
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_WHITE

# Slide 6: Packages
s6 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s6)
add_header(s6, "納品パッケージ一覧（TSH公式として即座に展開可能）")

delivs = [
    ("① 利用者目線 実践マニュアル（全20モジュール）", "HTML / Markdown / 印刷用PDF（現場NG操作・鉄則を完全網羅）"),
    ("② レスポンシブWebポータル一式", "スマホ・タブレット・PC全対応、高速インクリメンタル検索内蔵"),
    ("③ 実践「紙芝居」スライド動画（全10話）", "16:9 フルHD MP4、プロアナウンス調クリアナレーション音声付"),
    ("④ AIナレッジパック（NotebookLM / Gemini連携）", "Google NotebookLM / ChatGPT用統合ナレッジデータ一括提供")
]

for i, (title, desc) in enumerate(delivs):
    y = Inches(2.3 + i * 1.15)
    box = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), y, Inches(11.3), Inches(0.95))
    box.fill.solid()
    box.fill.fore_color.rgb = BG_CARD
    box.line.color.rgb = PRIMARY
    box.line.width = Pt(1.5)
    
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_MUTED

# Slide 7: Schedule
s7 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s7)
add_header(s7, "制作スケジュール（2〜3ヶ月で完工）")

schs = [
    ("Month 1（第1次納品）", "予約・フロント・マスター室編", "実践マニュアル ＋ 紙芝居動画4本 ＋ Webポータル初版"),
    ("Month 2（第2次納品）", "コンペ・顧客・会員管理・年会費編", "実践マニュアル ＋ 紙芝居動画3本 ＋ AIナレッジ統合"),
    ("Month 3（最終納品）", "売上日報・マスタ・AI連携完全版", "実践マニュアル ＋ 紙芝居動画3本 ＋ TSH公式リリース")
]

for i, (m, t, d) in enumerate(schs):
    x = Inches(1.0 + i * 3.8)
    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.4), Inches(3.6), Inches(3.8))
    card.fill.solid()
    card.fill.fore_color.rgb = BG_CARD
    card.line.color.rgb = PRIMARY
    card.line.width = Pt(2)
    
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"\n{m}\n"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    
    p2 = tf.add_paragraph()
    p2.text = f"\n{t}\n\n"
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    
    p3 = tf.add_paragraph()
    p3.text = d
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED

# Slide 8: Pricing
s8 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s8)
add_header(s8, "制作費用のご提案（一括受託・松竹梅プラン）")

plans = [
    ("【梅】スターター", "200 万円", "コア5モジュール\n紙芝居動画 3本\nWebポータル初版", BG_CARD, TEXT_WHITE),
    ("【竹】標準パッケージ（★推奨）", "350 万円", "主要10モジュール\n紙芝居動画 6本\nWeb ＋ AIナレッジパック\n（社内稟議最適ライン）", RGBColor(6, 78, 59), PRIMARY),
    ("【松】完全網羅フルパッケージ", "500 万円", "全20モジュール完全網羅\n紙芝居動画 10本全編\nGemini AIチャット内蔵\nTSH公式マニュアル完全刷新", BG_CARD, TEXT_WHITE),
]

for i, (name, price, desc, bg, title_col) in enumerate(plans):
    x = Inches(1.0 + i * 3.8)
    card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.3), Inches(3.6), Inches(4.2))
    card.fill.solid()
    card.fill.fore_color.rgb = bg
    card.line.color.rgb = title_col
    card.line.width = Pt(2)
    
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"\n{name}"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = title_col
    
    p2 = tf.add_paragraph()
    p2.text = f"{price} (税別)\n"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT
    
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_WHITE

# Slide 9: Next Step
s9 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s9)
add_header(s9, "次のステップ（まずは実機デモとNotebookLMをご体験ください）")

steps = [
    ("Step 1", "Google Drive共有フォルダから、Webポータル実機と紙芝居動画を閲覧"),
    ("Step 2", "オサケンさん共有のNotebookLMにて、AIとの自然言語・音声対話を実機テスト"),
    ("Step 3", "社内稟議用資料（Docsドラフト）をご活用いただき、役員会でのご検討・ご発注")
]

for i, (st, desc) in enumerate(steps):
    y = Inches(2.4 + i * 1.3)
    box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), y, Inches(11.3), Inches(1.0))
    box.fill.solid()
    box.fill.fore_color.rgb = BG_CARD
    box.line.color.rgb = PRIMARY
    box.line.width = Pt(1.5)
    
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"  {st}:  "
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(16)
    p2.font.color.rgb = TEXT_WHITE

# Slide 10: Closing
s10 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(s10)

tb = s10.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.3), Inches(2.0))
p = tb.text_frame.paragraphs[0]
p.text = "ご清聴ありがとうございました"
p.font.size = Pt(44)
p.font.bold = True
p.font.color.rgb = PRIMARY
p.alignment = PP_ALIGN.CENTER

tb2 = s10.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.3), Inches(1.5))
p = tb2.text_frame.paragraphs[0]
p.text = "SwingClub-CLOUD 次世代利用者ポータル ＆ 実践AIナレッジ\n津カントリー倶楽部 DXプロジェクト | 社長顧問 長田 賢一郎"
p.font.size = Pt(20)
p.font.color.rgb = TEXT_MUTED
p.alignment = PP_ALIGN.CENTER

prs.save(OUTPUT_PPTX)
print(f"Presentation saved successfully to: {OUTPUT_PPTX}")
